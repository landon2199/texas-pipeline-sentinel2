"""Gap 16: a gold-label sheet to measure the narrative reader (discovery/spill_reports.py), and its accuracy report.

  python gold_sheet.py --make [--n 30]   writes outputs/discovery/gold_labels.xlsx (and .csv): a random 30 of the 175
      right-of-way reports with the PHMSA narrative, and for each field (how found, soil removed, reached water,
      plants or crops mentioned) an empty column for a person's answer. The reader's own answers sit in a separate,
      hidden sheet so the person labels from the narrative alone.
  python gold_sheet.py --team-copy       writes the fill-in copy (no reader tab) to outputs/checks/gold_labels.xlsx
  python gold_sheet.py --score           reads the filled sheet (the outputs/checks copy if it exists) and writes outputs/discovery/gold_accuracy.md: per field
      the share the reader got right with a Wilson 95% interval, and a confusion table (good practice for thematic
      accuracy: Olofsson et al. 2014; Stehman and Foody 2019).
The sheet has no satellite results, so it is safe for anyone, including the photo checkers.
"""
import argparse
from pathlib import Path

import numpy as np
import pandas as pd

P = Path(r"C:\mydrive\Graduate School\Courses\GEOG_392\projects")
D = P / "outputs" / "discovery"
FIELDS = {"how_found": ["control room or alarm", "operator patrol or personnel", "landowner or public",
                        "third party contractor", "inspection or testing", "unknown"],
          "soil_removed": ["yes", "no", "unknown"], "reached_water": ["yes", "no", "unknown"],
          "plants_or_crops_mentioned": ["yes", "no", "unknown"]}


def make(n: int, seed: int):
    r = pd.read_parquet(D / "spill_reports.parquet")
    pick = r.sample(n=n, random_state=seed).reset_index(drop=True)
    sheet = pick[["REPORT_NUMBER", "LOCAL_DATETIME", "NARRATIVE"]].copy()
    for f, choices in FIELDS.items():
        sheet[f"label_{f}"] = ""
    sheet["notes"] = ""
    reader = pick[["REPORT_NUMBER"] + list(FIELDS)]
    help_ = pd.DataFrame({"field": list(FIELDS), "allowed answers": ["; ".join(c) for c in FIELDS.values()],
                          "how to answer": ["who first noticed the leak, from the narrative", "was contaminated soil dug up or hauled away",
                                            "did the oil reach a creek, pond, river or groundwater", "does the narrative mention plants, crops or vegetation"]})
    sheet.to_csv(D / "gold_labels.csv", index=False)
    try:
        with pd.ExcelWriter(D / "gold_labels.xlsx") as xw:
            help_.to_excel(xw, sheet_name="how to label", index=False)
            sheet.to_excel(xw, sheet_name="labels", index=False)
            reader.to_excel(xw, sheet_name="reader (hidden)", index=False)
        from openpyxl import load_workbook
        wb = load_workbook(D / "gold_labels.xlsx")
        wb["reader (hidden)"].sheet_state = "hidden"
        ws = wb["labels"]
        ws.column_dimensions["C"].width = 90
        for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
            for c in row:
                c.alignment = c.alignment.copy(wrap_text=True, vertical="top")
        wb.save(D / "gold_labels.xlsx")
        print(f"wrote {D / 'gold_labels.xlsx'} and .csv ({n} reports)")
    except ImportError:
        print(f"openpyxl is missing; wrote {D / 'gold_labels.csv'} only")


def wilson(k: int, n: int):
    if n == 0:
        return np.nan, np.nan
    z, p = 1.96, k / n
    c = (p + z * z / (2 * n)) / (1 + z * z / n)
    h = z * np.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / (1 + z * z / n)
    return c - h, c + h


def team_copy():
    """The sheet the team fills in: outputs/checks/gold_labels.xlsx, with the help tab and the labels tab only, so the
    reader's answers can't be unhidden. Scoring reads the reader's answers from spill_reports.parquet, not the sheet."""
    out = P / "outputs" / "checks" / "gold_labels.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(out) as xw:
        for name in ("how to label", "labels"):
            pd.read_excel(D / "gold_labels.xlsx", sheet_name=name, dtype=str).to_excel(xw, sheet_name=name, index=False)
    from openpyxl import load_workbook
    wb = load_workbook(out)
    wb["labels"].column_dimensions["C"].width = 90
    for row in wb["labels"].iter_rows(min_row=2, min_col=3, max_col=3):
        for c in row:
            c.alignment = c.alignment.copy(wrap_text=True, vertical="top")
    wb.save(out)
    print(f"wrote {out} (no reader tab)")


def score():
    team = P / "outputs" / "checks" / "gold_labels.xlsx"          # the team's filled copy, if they used it
    lab = pd.read_excel(team if team.exists() else D / "gold_labels.xlsx", sheet_name="labels", dtype=str)
    r = pd.read_parquet(D / "spill_reports.parquet", columns=["REPORT_NUMBER"] + list(FIELDS)).astype({"REPORT_NUMBER": str})
    m = lab.astype({"REPORT_NUMBER": str}).merge(r, on="REPORT_NUMBER", how="left")
    lines = [f"# How accurate is the narrative reader? ({pd.Timestamp.today():%Y-%m-%d})", "",
             "| Field | Labeled | Reader right | Accuracy [Wilson 95% CI] |", "|---|---|---|---|"]
    tables = []
    for f in FIELDS:
        g = m[m[f"label_{f}"].fillna("").str.strip() != ""]
        right = (g[f"label_{f}"].str.strip().str.lower() == g[f].astype(str).str.lower())
        lo, hi = wilson(int(right.sum()), len(g))
        lines.append(f"| {f} | {len(g)} | {int(right.sum())} | {right.mean():.0%} [{lo:.0%}, {hi:.0%}] |" if len(g) else f"| {f} | 0 | - | - |")
        if len(g):
            tables.append(f"\n**{f}** (rows: the person's label; columns: the reader's answer)\n\n" +
                          pd.crosstab(g[f"label_{f}"].str.strip().str.lower(), g[f].astype(str).str.lower()).to_markdown())
    (D / "gold_accuracy.md").write_text("\n".join(lines) + "\n" + "\n".join(tables) + "\n", encoding="utf-8")
    print("\n".join(lines))


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--make", action="store_true")
    ap.add_argument("--score", action="store_true")
    ap.add_argument("--team-copy", action="store_true", help="write the fill-in copy to outputs/checks")
    ap.add_argument("--n", type=int, default=30)
    ap.add_argument("--seed", type=int, default=392)
    a = ap.parse_args()
    make(a.n, a.seed) if a.make else score() if a.score else team_copy() if a.team_copy else ap.print_help()
